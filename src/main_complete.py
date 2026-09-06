import flet as ft
import sqlite3
import json
from datetime import datetime, date, timedelta
from pathlib import Path
import shutil
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# الإعدادات العامة
# ============================================================

APP_NAME_AR = "المركز الشخصي الذكي"
APP_NAME_EN = "Smart Personal Hub"
BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "hub.db"
ASSETS_DIR = BASE_DIR / "assets"
CHARTS_DIR = ASSETS_DIR / "charts"
COVERS_DIR = ASSETS_DIR / "covers"
EXPORTS_DIR = BASE_DIR / "exports"

for dir_path in [ASSETS_DIR, CHARTS_DIR, COVERS_DIR, EXPORTS_DIR]:
    dir_path.mkdir(exist_ok=True)

# ============================================================
# قاعدة البيانات - محسّنة
# ============================================================

class Database:
    """إدارة قاعدة البيانات SQLite"""

    def __init__(self, page):
        self.page = page
        self.conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self.create_tables()

    def create_tables(self):
        """إنشاء جميع الجداول"""
        cursor = self.conn.cursor()

        # جدول المهام
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                description TEXT,
                task_date TEXT,
                priority TEXT DEFAULT 'normal',
                completed INTEGER DEFAULT 0,
                created_at TEXT
            )
        """)

        # جدول المصاريف
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                type TEXT NOT NULL,
                amount REAL NOT NULL,
                category TEXT,
                expense_date TEXT,
                note TEXT,
                created_at TEXT
            )
        """)

        # جدول الكتب
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS books (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                author TEXT,
                cover_path TEXT,
                progress INTEGER DEFAULT 0,
                rating INTEGER,
                start_date TEXT,
                end_date TEXT,
                created_at TEXT
            )
        """)

        # جدول كلمات السر
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS passwords (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                site_name TEXT NOT NULL,
                email TEXT,
                password TEXT,
                note TEXT,
                created_at TEXT
            )
        """)

        # جدول الوصفات
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS recipes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                image_path TEXT,
                ingredients TEXT,
                steps TEXT,
                cook_time INTEGER,
                created_at TEXT
            )
        """)

        # جدول الجدول الدراسي
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS schedule (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                date TEXT,
                type TEXT,
                description TEXT,
                created_at TEXT
            )
        """)

        # جدول الفواتير
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS invoices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                invoice_number TEXT UNIQUE,
                date TEXT,
                items TEXT,
                total REAL,
                created_at TEXT
            )
        """)

        # جدول تقدم القرآن
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS quran_progress (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                surah_number INTEGER,
                ayah_number INTEGER,
                progress REAL DEFAULT 0,
                updated_at TEXT
            )
        """)

        self.conn.commit()

    # دوال مشتركة
    def add_item(self, table, **kwargs):
        """إضافة عنصر عام"""
        columns = list(kwargs.keys())
        values = list(kwargs.values())
        placeholders = ",".join(["?" for _ in columns])
        query = f"INSERT INTO {table} ({','.join(columns)}) VALUES ({placeholders})"
        self.conn.execute(query, values)
        self.conn.commit()

    def get_items(self, table, where=None):
        """جلب عناصر"""
        query = f"SELECT * FROM {table}"
        if where:
            query += f" WHERE {where}"
        return self.conn.execute(query).fetchall()

    def delete_item(self, table, item_id):
        """حذف عنصر"""
        self.conn.execute(f"DELETE FROM {table} WHERE id = ?", (item_id,))
        self.conn.commit()

# ============================================================
# أدوات واجهة المستخدم - محسّنة
# ============================================================

class UI:
    """أدوات تصميم موحدة"""

    def __init__(self, page, app):
        self.page = page
        self.app = app

    def card(self, content, padding=18, shadow=True):
        """بطاقة حديثة"""
        return ft.Container(
            content=content,
            padding=padding,
            border_radius=15,
            bgcolor=ft.Colors.SURFACE_CONTAINER,
            border=ft.border.all(1, ft.Colors.OUTLINE_VARIANT),
            shadow=ft.BoxShadow(
                blur_radius=15,
                spread_radius=0,
                offset=ft.Offset(0, 5),
                color=ft.Colors.BLACK26
            ) if shadow else None,
            animate_scale=ft.animation.Animation(300, ft.AnimationCurve.EASE_OUT),
        )

    def button(self, text, icon, on_click):
        """زر موحد"""
        return ft.ElevatedButton(
            text=text,
            icon=icon,
            on_click=on_click,
            style=ft.ButtonStyle(
                shape=ft.RoundedRectangleBorder(radius=12),
                padding=ft.padding.symmetric(horizontal=18, vertical=14)
            )
        )

    def show_snack(self, message):
        """رسالة"""
        self.page.snack_bar = ft.SnackBar(
            content=ft.Text(message),
            open=True
        )
        self.page.update()

    def input_field(self, label, icon, multiline=False, keyboard_type=ft.KeyboardType.TEXT):
        """حقل إدخال"""
        return ft.TextField(
            label=label,
            prefix_icon=icon,
            border_radius=12,
            expand=True,
            multiline=multiline,
            min_lines=2 if multiline else 1,
            keyboard_type=keyboard_type
        )

# ============================================================
# القسم 1: منظم المهام (موجود بالفعل)
# ============================================================

class TaskManager:
    def __init__(self, page, db, app):
        self.page = page
        self.db = db
        self.app = app
        self.ui = UI(page, app)
        self.filter_mode = "all"
        self.title_field = self.ui.input_field("اسم المهمة", ft.Icons.TASK_ALT)
        self.description_field = self.ui.input_field("الوصف", ft.Icons.DESCRIPTION, multiline=True)
        self.date_field = ft.TextField(
            label="التاريخ",
            prefix_icon=ft.Icons.CALENDAR_MONTH,
            value=date.today().isoformat(),
            read_only=True,
            border_radius=12,
            expand=True
        )
        self.priority_dropdown = ft.Dropdown(
            label="الأولوية",
            value="normal",
            options=[
                ft.dropdown.Option("urgent", "عاجل"),
                ft.dropdown.Option("important", "مهم"),
                ft.dropdown.Option("normal", "عادي")
            ],
            border_radius=12,
            expand=True
        )
        self.task_list = ft.Column(spacing=10, scroll=ft.ScrollMode.AUTO, expand=True)
        self.stats_row = ft.Row(spacing=12, wrap=True)

    def choose_date(self, e):
        def on_date_change(ev):
            if ev.control.value:
                self.date_field.value = ev.control.value.strftime("%Y-%m-%d")
            self.page.update()
        picker = ft.DatePicker(on_change=on_date_change)
        self.page.overlay.append(picker)
        picker.open = True
        self.page.update()

    def add_task(self, e):
        title = self.title_field.value.strip()
        if not title:
            self.ui.show_snack("يرجى كتابة اسم المهمة")
            return
        self.db.add_item(
            "tasks",
            title=title,
            description=self.description_field.value.strip(),
            task_date=self.date_field.value,
            priority=self.priority_dropdown.value,
            created_at=datetime.now().isoformat()
        )
        self.title_field.value = ""
        self.description_field.value = ""
        self.refresh()
        self.ui.show_snack("تمت إضافة المهمة")

    def set_filter(self, mode):
        self.filter_mode = mode
        self.refresh()

    def delete_task(self, task_id):
        self.db.delete_item("tasks", task_id)
        self.refresh()
        self.ui.show_snack("تم حذف المهمة")

    def refresh(self):
        self.task_list.controls.clear()
        tasks = self.db.get_items("tasks")
        if not tasks:
            self.task_list.controls.append(
                ft.Container(
                    content=ft.Column(
                        [ft.Icon(ft.Icons.TASK_ALT, size=55, color=ft.Colors.ON_SURFACE_VARIANT),
                         ft.Text("لا توجد مهام حالياً", size=16)],
                        horizontal_alignment=ft.CrossAxisAlignment.CENTER
                    ),
                    padding=40,
                    alignment=ft.alignment.center
                )
            )
        else:
            for task in tasks:
                self.task_list.controls.append(self.ui.card(
                    ft.Column([
                        ft.Row([
                            ft.Text(task["title"], size=17, weight=ft.FontWeight.BOLD, expand=True),
                            ft.IconButton(ft.Icons.DELETE_OUTLINE, on_click=lambda e, tid=task["id"]: self.delete_task(tid))
                        ]),
                        ft.Text(f"الأولوية: {task['priority']} | التاريخ: {task['task_date']}", size=12, color=ft.Colors.ON_SURFACE_VARIANT)
                    ], spacing=5)
                ))
        self.page.update()

    def build(self):
        self.refresh()
        return ft.Column([
            ft.Text("منظم المهام", size=30, weight=ft.FontWeight.BOLD),
            self.ui.card(ft.Column([
                ft.Text("إضافة مهمة", size=20, weight=ft.FontWeight.BOLD),
                self.title_field,
                self.description_field,
                ft.Row([self.date_field, ft.IconButton(ft.Icons.CALENDAR_MONTH, on_click=self.choose_date), self.priority_dropdown]),
                self.ui.button("إضافة", ft.Icons.ADD, self.add_task)
            ], spacing=12)),
            self.ui.card(ft.Column([
                ft.Text("قائمة المهام", size=20, weight=ft.FontWeight.BOLD),
                self.task_list
            ], spacing=12, expand=True))
        ], spacing=15, expand=True, scroll=ft.ScrollMode.AUTO)

# ============================================================
# القسم 2: حاسبة المصاريف (موجود بالفعل)
# ============================================================

class ExpenseManager:
    def __init__(self, page, db, app):
        self.page = page
        self.db = db
        self.app = app
        self.ui = UI(page, app)
        self.selected_month = date.today().strftime("%Y-%m")
        self.type_dropdown = ft.Dropdown(
            label="النوع",
            value="expense",
            options=[ft.dropdown.Option("expense", "مصروف"), ft.dropdown.Option("income", "دخل")],
            border_radius=12,
            expand=True
        )
        self.amount_field = self.ui.input_field("المبلغ", ft.Icons.ATTACH_MONEY, keyboard_type=ft.KeyboardType.NUMBER)
        self.category_field = self.ui.input_field("التصنيف", ft.Icons.CATEGORY)
        self.date_field = ft.TextField(
            label="التاريخ",
            value=date.today().isoformat(),
            read_only=True,
            border_radius=12,
            expand=True
        )
        self.expense_list = ft.Column(spacing=8, scroll=ft.ScrollMode.AUTO, expand=True)
        self.stats_row = ft.Row(spacing=12, wrap=True)
        self.chart_image = ft.Image(visible=False, fit=ft.ImageFit.contain, height=300)

    def add_transaction(self, e):
        try:
            amount = float(self.amount_field.value)
        except:
            self.ui.show_snack("يرجى إدخال مبلغ صحيح")
            return
        self.db.add_item(
            "expenses",
            type=self.type_dropdown.value,
            amount=amount,
            category=self.category_field.value or "عام",
            expense_date=self.date_field.value,
            created_at=datetime.now().isoformat()
        )
        self.amount_field.value = ""
        self.category_field.value = ""
        self.refresh()
        self.ui.show_snack("تم حفظ العملية")

    def generate_chart(self):
        expenses = self.db.get_items("expenses", f"strftime('%Y-%m', expense_date) = '{self.selected_month}' AND type = 'expense'")
        if not expenses:
            self.chart_image.visible = False
            return
        
        categories = {}
        for exp in expenses:
            cat = exp["category"] or "عام"
            categories[cat] = categories.get(cat, 0) + exp["amount"]
        
        plt.figure(figsize=(6, 4), dpi=120)
        plt.pie(categories.values(), labels=categories.keys(), autopct="%1.1f%%", startangle=90)
        plt.title(f"المصاريف - {self.selected_month}")
        plt.tight_layout()
        
        chart_path = CHARTS_DIR / "expenses.png"
        plt.savefig(chart_path, bbox_inches="tight")
        plt.close()
        
        self.chart_image.src = str(chart_path)
        self.chart_image.visible = True

    def refresh(self):
        self.expense_list.controls.clear()
        expenses = self.db.get_items("expenses", f"strftime('%Y-%m', expense_date) = '{self.selected_month}'")
        
        if not expenses:
            self.expense_list.controls.append(ft.Container(
                content=ft.Text("لا توجد عمليات مالية لهذا الشهر"),
                padding=40,
                alignment=ft.alignment.center
            ))
        else:
            for exp in expenses:
                self.expense_list.controls.append(self.ui.card(
                    ft.Row([
                        ft.Icon(ft.Icons.TRENDING_UP if exp["type"] == "income" else ft.Icons.TRENDING_DOWN, size=25),
                        ft.Column([
                            ft.Text(exp["category"], size=14, weight=ft.FontWeight.BOLD),
                            ft.Text(exp["expense_date"], size=12, color=ft.Colors.ON_SURFACE_VARIANT)
                        ], expand=True),
                        ft.Text(f"{exp['amount']:.2f}", size=16, weight=ft.FontWeight.BOLD),
                        ft.IconButton(ft.Icons.DELETE, on_click=lambda e, eid=exp["id"]: self.delete_expense(eid))
                    ])
                ))
        
        self.generate_chart()
        self.page.update()

    def delete_expense(self, expense_id):
        self.db.delete_item("expenses", expense_id)
        self.refresh()

    def build(self):
        self.refresh()
        return ft.Column([
            ft.Text("حاسبة المصاريف", size=30, weight=ft.FontWeight.BOLD),
            self.ui.card(ft.Column([
                ft.Text("إضافة عملية", size=20, weight=ft.FontWeight.BOLD),
                ft.Row([self.type_dropdown, self.amount_field]),
                ft.Row([self.category_field, self.date_field]),
                self.ui.button("حفظ", ft.Icons.SAVE, self.add_transaction)
            ], spacing=12)),
            self.ui.card(ft.Column([ft.Text("الرسم البياني", size=20, weight=ft.FontWeight.BOLD), self.chart_image], horizontal_alignment=ft.CrossAxisAlignment.CENTER)),
            self.ui.card(ft.Column([ft.Text("العمليات المالية", size=20, weight=ft.FontWeight.BOLD), self.expense_list], expand=True))
        ], spacing=15, expand=True, scroll=ft.ScrollMode.AUTO)

# ============================================================
# الأقسام المتبقية - Placeholders محسّنة
# ============================================================

class PlaceholderSection:
    def __init__(self, page, app, title_ar, title_en, icon, color=ft.Colors.PRIMARY):
        self.page = page
        self.app = app
        self.title_ar = title_ar
        self.title_en = title_en
        self.icon = icon
        self.color = color

    def build(self):
        title = self.title_ar if self.app.language == "ar" else self.title_en
        return ft.Container(
            content=ft.Column([
                ft.Icon(self.icon, size=80, color=self.color),
                ft.Text(title, size=30, weight=ft.FontWeight.BOLD),
                ft.Text("هذا القسم في مرحلة التطوير", size=15, color=ft.Colors.ON_SURFACE_VARIANT, text_align=ft.TextAlign.CENTER),
                ft.Container(height=20),
                ft.ElevatedButton("قريباً", icon=ft.Icons.STAR, disabled=True)
            ], horizontal_alignment=ft.CrossAxisAlignment.CENTER, alignment=ft.MainAxisAlignment.CENTER, spacing=15),
            expand=True,
            alignment=ft.alignment.center
        )

# ============================================================
# التطبيق الرئيسي - محسّن
# ============================================================

class SmartPersonalHub:
    def __init__(self, page):
        self.page = page
        self.language = getattr(self.page.session, "language", "ar")
        self.theme_mode = getattr(self.page.session, "theme_mode", "dark")
        
        self.db = Database(page)
        self.task_manager = TaskManager(page, self.db, self)
        self.expense_manager = ExpenseManager(page, self.db, self)
        self.current_section = 0

    def configure_page(self):
        self.page.title = APP_NAME_AR if self.language == "ar" else APP_NAME_EN
        self.page.padding = 0
        self.page.theme = ft.Theme(font_family="Cairo")
        self.page.dark_theme = ft.Theme(font_family="Cairo")
        self.page.theme_mode = ft.ThemeMode.DARK if self.theme_mode == "dark" else ft.ThemeMode.LIGHT
        self.page.rtl = self.language == "ar"

    def toggle_theme(self, e):
        self.theme_mode = "light" if self.theme_mode == "dark" else "dark"
        self.page.theme_mode = ft.ThemeMode.LIGHT if self.theme_mode == "light" else ft.ThemeMode.DARK
        self.page.session.theme_mode = self.theme_mode
        self.page.update()

    def toggle_language(self, e):
        self.language = "en" if self.language == "ar" else "ar"
        self.page.session.language = self.language
        self.configure_page()
        self.build()

    def build_header(self):
        title = APP_NAME_AR if self.language == "ar" else APP_NAME_EN
        subtitle = "إدارة حياتك من مكان واحد" if self.language == "ar" else "Manage your life from one place"
        theme_icon = ft.Icons.LIGHT_MODE if self.theme_mode == "dark" else ft.Icons.DARK_MODE
        lang_text = "English" if self.language == "ar" else "العربية"

        return ft.Container(
            content=ft.Row([
                ft.Column([
                    ft.Text(title, size=24, weight=ft.FontWeight.BOLD),
                    ft.Text(subtitle, size=12, color=ft.Colors.ON_SURFACE_VARIANT)
                ], spacing=2, expand=True),
                ft.IconButton(theme_icon, on_click=self.toggle_theme),
                ft.OutlinedButton(lang_text, icon=ft.Icons.LANGUAGE, on_click=self.toggle_language)
            ], vertical_alignment=ft.CrossAxisAlignment.CENTER),
            padding=ft.padding.symmetric(horizontal=20, vertical=12),
            border=ft.border.only(bottom=ft.BorderSide(1, ft.Colors.OUTLINE_VARIANT))
        )

    def navigation_items(self):
        labels_ar = ["المهام", "المصاريف", "الكتب", "كلمات السر", "الوحدات", "الجدول", "الوصفات", "الفواتير", "القرآن", "الملفات"]
        labels_en = ["Tasks", "Expenses", "Books", "Passwords", "Units", "Schedule", "Recipes", "Invoices", "Quran", "Files"]
        labels = labels_ar if self.language == "ar" else labels_en
        
        icons = [
            (ft.Icons.TASK_ALT, ft.Icons.TASK_ALT),
            (ft.Icons.ACCOUNT_BALANCE_WALLET_OUTLINED, ft.Icons.ACCOUNT_BALANCE_WALLET),
            (ft.Icons.BOOK_OUTLINED, ft.Icons.BOOK),
            (ft.Icons.LOCK_OUTLINE, ft.Icons.LOCK),
            (ft.Icons.SWAP_HORIZ_OUTLINED, ft.Icons.SWAP_HORIZ),
            (ft.Icons.CALENDAR_MONTH_OUTLINED, ft.Icons.CALENDAR_MONTH),
            (ft.Icons.RESTAURANT_MENU_OUTLINED, ft.Icons.RESTAURANT_MENU),
            (ft.Icons.RECEIPT_LONG_OUTLINED, ft.Icons.RECEIPT_LONG),
            (ft.Icons.MENU_BOOK_OUTLINED, ft.Icons.MENU_BOOK),
            (ft.Icons.FOLDER_OUTLINED, ft.Icons.FOLDER)
        ]

        return [ft.NavigationRailDestination(icon=icons[i][0], selected_icon=icons[i][1], label=labels[i]) for i in range(10)]

    def change_section(self, e):
        self.current_section = e.control.selected_index
        self.content_area.controls.clear()

        if self.current_section == 0:
            self.content_area.controls.append(self.task_manager.build())
        elif self.current_section == 1:
            self.content_area.controls.append(self.expense_manager.build())
        else:
            sections_data = [
                ("مكتبة الكتب", "Books Library", ft.Icons.BOOK, ft.Colors.BLUE),
                ("مدير كلمات السر", "Password Manager", ft.Icons.LOCK, ft.Colors.RED),
                ("محول الوحدات", "Unit Converter", ft.Icons.SWAP_HORIZ, ft.Colors.ORANGE),
                ("الجدول الدراسي", "Study Schedule", ft.Icons.CALENDAR_MONTH, ft.Colors.GREEN),
                ("مخزن الوصفات", "Recipe Manager", ft.Icons.RESTAURANT_MENU, ft.Colors.PURPLE),
                ("برنامج الفواتير", "Invoice Manager", ft.Icons.RECEIPT_LONG, ft.Colors.YELLOW),
                ("تطبيق القرآن", "Quran", ft.Icons.MENU_BOOK, ft.Colors.CYAN),
                ("منظم الملفات", "File Organizer", ft.Icons.FOLDER, ft.Colors.AMBER),
            ]
            
            data = sections_data[self.current_section - 2]
            self.content_area.controls.append(PlaceholderSection(self.page, self, data[0], data[1], data[2], data[3]).build())

        self.page.update()

    def build(self):
        self.configure_page()
        self.page.controls.clear()

        header = self.build_header()
        self.content_area = ft.Column(expand=True, spacing=0)

        if self.current_section == 0:
            self.content_area.controls.append(self.task_manager.build())
        else:
            self.content_area.controls.append(self.expense_manager.build())

        navigation = ft.NavigationRail(
            selected_index=self.current_section,
            label_type=ft.NavigationRailLabelType.ALL,
            destinations=self.navigation_items(),
            on_change=self.change_section,
            extended=False,
            min_width=85
        )

        main_row = ft.Row([
            navigation,
            ft.VerticalDivider(width=1),
            ft.Container(content=self.content_area, expand=True, padding=20)
        ], expand=True, spacing=0)

        self.page.add(header, main_row)
        self.page.update()

def main(page: ft.Page):
    app = SmartPersonalHub(page)
    app.build()

ft.run(target=main)
